from __future__ import annotations

import logging
import os
from typing import Dict, List, Optional, Type

from rfq.interfaces import DocumentParser

logger = logging.getLogger(__name__)


class PdfParser:
    """Extract text from PDF files using PyPDF2. For scanned PDFs, AI will handle extraction."""

    @staticmethod
    def supported_extensions() -> List[str]:
        return ['.pdf']

    @staticmethod
    def extract_text(file_path: str) -> str:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            text = '\n'.join(
                page.extract_text() or '' for page in reader.pages
            )
            logger.info(
                'PDF extraction: %d chars extracted from %s (%d pages)',
                len(text), file_path, len(reader.pages),
            )
            return text
        except Exception as exc:
            logger.error('PDF extraction error [%s]: %s', file_path, exc)
            return ''


class DocxParser:
    """Extract text from DOCX files using python-docx."""

    @staticmethod
    def supported_extensions() -> List[str]:
        return ['.docx', '.doc']

    @staticmethod
    def extract_text(file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text for cell in row.cells)
                    text += '\n' + row_text
            logger.info(
                'DOCX extraction: %d chars extracted from %s',
                len(text), file_path,
            )
            return text
        except Exception as exc:
            logger.error('DOCX extraction error [%s]: %s', file_path, exc)
            return ''


# ---------------------------------------------------------------------------
# Registry / factory
# ---------------------------------------------------------------------------

_PARSERS: Dict[str, Type[DocumentParser]] = {
    ext: cls
    for cls in [PdfParser, DocxParser]
    for ext in cls.supported_extensions()
}


def get_parser(file_path: str) -> Optional[DocumentParser]:
    """
    Return the appropriate parser for *file_path* based on its extension.
    """
    _, ext = os.path.splitext(file_path)
    cls = _PARSERS.get(ext.lower())
    if cls is None:
        logger.warning('No parser for extension %s', ext)
        return None
    return cls()


def extract_text(file_path: str) -> str:
    """Convenience: parse *file_path* and return its text content."""
    parser = get_parser(file_path)
    if parser is None:
        return ''
    return parser.extract_text(file_path)
